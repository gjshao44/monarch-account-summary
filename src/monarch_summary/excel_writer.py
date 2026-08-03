"""Write ledger rows into the Data_AZS / Data_SSP tabs of the budget workbook.

Only columns A-G (Mo, Date, Account, Description, Amount, Label, Notes) of
those two sheets are ever touched, and every other sheet/formula in the
workbook is left completely alone -- the Monthly_AZS / Monthly_SSP /
Monthy_Comb tabs read from these two sheets and do the rest of the math
themselves via SUMIFS against column A's month number.

Column A is normally pre-filled with `=MONTH(B..)` in the template and is
left alone when present. However the observed template has gaps (e.g.
Data_SSP's column A formula only starts at row 8, not row 6) where a written
row would otherwise get silently excluded from every Monthly_* rollup despite
having a correct Date/Amount -- so any row this module writes into gets its
column A formula filled in if it's missing, without disturbing rows it
doesn't touch.
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import openpyxl

from monarch_summary.categorize import LedgerRow

SHEET_BY_USER = {"AZS": "Data_AZS", "SSP": "Data_SSP"}
DATA_START_ROW = 6
COL_MONTH = 1
COL_DATE = 2
COL_ACCOUNT = 3
COL_DESCRIPTION = 4
COL_AMOUNT = 5
COL_LABEL = 6
COL_NOTES = 7
LAST_DATA_COL = 7  # G


def write_ledger(
    rows: list[LedgerRow],
    workbook_path: str | Path,
    output_path: str | Path | None = None,
    in_place: bool = False,
) -> Path:
    """Write rows into a copy of workbook_path's Data_AZS/Data_SSP tabs.

    Returns the path the workbook was saved to. Re-running is idempotent:
    the existing data region (rows DATA_START_ROW..sheet.max_row, cols B-G)
    is cleared before writing, since a Monarch CSV export is a full-history
    dump rather than an incremental delta.
    """
    workbook_path = Path(workbook_path)
    wb = openpyxl.load_workbook(workbook_path)

    by_user: dict[str, list[LedgerRow]] = {"AZS": [], "SSP": []}
    for row in rows:
        if row.user not in by_user:
            raise ValueError(f"Unknown user code {row.user!r} on ledger row: {row}")
        by_user[row.user].append(row)

    for user, sheet_name in SHEET_BY_USER.items():
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Workbook is missing expected sheet {sheet_name!r}")
        ws = wb[sheet_name]
        _clear_data_region(ws)

        user_rows = sorted(by_user[user], key=lambda r: r.date)
        capacity = ws.max_row - DATA_START_ROW + 1
        if len(user_rows) > capacity:
            raise ValueError(
                f"{sheet_name} only has room for {capacity} transaction rows "
                f"(rows {DATA_START_ROW}-{ws.max_row}), but {len(user_rows)} "
                f"were routed to {user}. Extend the sheet's pre-filled rows "
                "in the template first."
            )

        for i, row in enumerate(user_rows):
            r = DATA_START_ROW + i
            month_cell = ws.cell(row=r, column=COL_MONTH)
            if month_cell.value in (None, ""):
                month_cell.value = f"=MONTH(B{r})"
            # NOTE: ws.cell(..., value=None) is a no-op in openpyxl (it means
            # "don't set"), not "clear" -- so cell.value must be assigned
            # directly to correctly write a None label.
            ws.cell(row=r, column=COL_DATE).value = _as_datetime(row.date)
            ws.cell(row=r, column=COL_ACCOUNT).value = row.account
            ws.cell(row=r, column=COL_DESCRIPTION).value = row.description
            ws.cell(row=r, column=COL_AMOUNT).value = row.amount
            ws.cell(row=r, column=COL_LABEL).value = row.label  # None -> blank cell
            ws.cell(row=r, column=COL_NOTES).value = row.notes

    if in_place:
        dest = workbook_path
    elif output_path is not None:
        dest = Path(output_path)
    else:
        dest = Path("expense_data/output") / f"{workbook_path.stem}_synced{workbook_path.suffix}"

    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    return dest


def _clear_data_region(ws) -> None:
    for r in range(DATA_START_ROW, ws.max_row + 1):
        for c in range(COL_DATE, LAST_DATA_COL + 1):
            ws.cell(row=r, column=c).value = None


def _as_datetime(d: date) -> datetime:
    if isinstance(d, datetime):
        return d
    return datetime(d.year, d.month, d.day)
