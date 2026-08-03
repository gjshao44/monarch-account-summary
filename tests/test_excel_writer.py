from datetime import date

import openpyxl
import pytest

from monarch_summary.categorize import LedgerRow
from monarch_summary.excel_writer import DATA_START_ROW, write_ledger


def make_row(user, day, amount, label="Restaurants", desc="M"):
    return LedgerRow(
        user=user,
        date=date(2026, 1, day),
        account="Acct",
        description=desc,
        amount=amount,
        label=label,
        notes="n",
    )


def test_write_ledger_basic(tiny_workbook_path, tmp_path):
    rows = [make_row("AZS", 2, -10.0), make_row("SSP", 3, -20.0)]
    dest = write_ledger(rows, tiny_workbook_path, output_path=tmp_path / "out.xlsx")

    wb = openpyxl.load_workbook(dest)
    azs, ssp = wb["Data_AZS"], wb["Data_SSP"]

    assert azs.cell(row=DATA_START_ROW, column=5).value == -10.0
    assert azs.cell(row=DATA_START_ROW, column=1).value == f"=MONTH(B{DATA_START_ROW})"
    assert ssp.cell(row=DATA_START_ROW, column=5).value == -20.0
    assert azs.cell(row=DATA_START_ROW + 1, column=2).value is None


def test_write_ledger_idempotent_rerun(tiny_workbook_path, tmp_path):
    out = tmp_path / "out.xlsx"
    write_ledger(
        [make_row("AZS", 2, -10.0), make_row("AZS", 3, -20.0)],
        tiny_workbook_path,
        output_path=out,
    )

    dest = write_ledger([make_row("AZS", 5, -99.0)], out, output_path=out)

    wb = openpyxl.load_workbook(dest)
    azs = wb["Data_AZS"]
    assert azs.cell(row=DATA_START_ROW, column=5).value == -99.0
    assert azs.cell(row=DATA_START_ROW + 1, column=5).value is None


def test_write_ledger_capacity_exceeded(tiny_workbook_path, tmp_path):
    rows = [make_row("AZS", d, -1.0) for d in range(2, 12)]  # tiny fixture capacity is 5
    with pytest.raises(ValueError, match="only has room"):
        write_ledger(rows, tiny_workbook_path, output_path=tmp_path / "out.xlsx")


def test_write_ledger_sorts_by_date(tiny_workbook_path, tmp_path):
    rows = [make_row("AZS", 5, -1.0), make_row("AZS", 2, -2.0)]
    dest = write_ledger(rows, tiny_workbook_path, output_path=tmp_path / "out.xlsx")

    wb = openpyxl.load_workbook(dest)
    azs = wb["Data_AZS"]
    assert azs.cell(row=DATA_START_ROW, column=5).value == -2.0
    assert azs.cell(row=DATA_START_ROW + 1, column=5).value == -1.0


def test_self_heals_missing_column_a_formula(tiny_workbook_path, tmp_path):
    # Mirrors a real gap found in the actual template: Data_SSP's column A
    # =MONTH(B..) formula only starts partway down, leaving early rows blank.
    wb = openpyxl.load_workbook(tiny_workbook_path)
    wb["Data_SSP"].cell(row=DATA_START_ROW, column=1).value = None
    wb.save(tiny_workbook_path)

    dest = write_ledger(
        [make_row("SSP", 2, -5.0)], tiny_workbook_path, output_path=tmp_path / "out.xlsx"
    )

    wb = openpyxl.load_workbook(dest)
    assert wb["Data_SSP"].cell(row=DATA_START_ROW, column=1).value == f"=MONTH(B{DATA_START_ROW})"


def test_null_label_written_as_blank_cell(tiny_workbook_path, tmp_path):
    dest = write_ledger(
        [make_row("AZS", 2, -1.0, label=None)], tiny_workbook_path, output_path=tmp_path / "out.xlsx"
    )
    wb = openpyxl.load_workbook(dest)
    assert wb["Data_AZS"].cell(row=DATA_START_ROW, column=6).value is None
