import openpyxl

from conftest import build_tiny_workbook
from monarch_summary.pipeline import sync_from_csv


def test_sync_from_csv_end_to_end(fixtures_dir, tmp_path):
    workbook_path = tmp_path / "workbook.xlsx"
    build_tiny_workbook(workbook_path, capacity=10)

    outcome = sync_from_csv(
        transactions_csv_path=fixtures_dir / "sample_transactions.csv",
        workbook_path=workbook_path,
        owners_config_path=fixtures_dir / "owners.test.yaml",
        categories_config_path=fixtures_dir / "categories.test.yaml",
        output_path=tmp_path / "out.xlsx",
    )

    assert outcome.parse_result.total_rows == 10
    assert outcome.parse_result.skipped_blank_amount == 1

    assert set(outcome.summaries) == {"AZS", "SSP", "Combined"}
    assert outcome.summaries["AZS"]["row_count"] == 7  # 6 direct + 1 shared-split half
    assert outcome.summaries["SSP"]["row_count"] == 3  # 2 direct + 1 shared-split half
    assert outcome.summaries["Combined"]["row_count"] == 10

    # Restaurants(-12.50) + Public Transport(-3.00) + Business->Misc(-15.00)
    # + shared Restaurants half(-30.00); Transfer(-500) and the unmapped
    # "Weird Category" row are excluded from spend entirely.
    assert outcome.summaries["AZS"]["spend"] == 60.5
    assert outcome.cat_result.unmapped_categories == {"Weird Category": 1}

    wb = openpyxl.load_workbook(outcome.dest)
    assert "Data_AZS" in wb.sheetnames and "Data_SSP" in wb.sheetnames
