"""Shared pytest fixtures. Everything here uses synthetic data only."""

from pathlib import Path

import openpyxl
import pytest

DATA_SHEETS = ["Data_AZS", "Data_SSP"]


def build_tiny_workbook(path: Path, capacity: int = 5) -> None:
    """Build a minimal workbook mimicking the real template's Data_AZS/Data_SSP
    layout: header at row 4, blank spacer at row 5, pre-filled `=MONTH(B..)`
    formulas in column A from row 6 through row 5+capacity.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for sheet_name in DATA_SHEETS:
        ws = wb.create_sheet(sheet_name)
        ws["A1"] = f"{sheet_name} Account Data"
        ws.append(["Mo", "Date", "Account", "Description", "Amount ($)", "Label", "Notes"])  # would land row 2 without spacer below
        # Force header onto row 4 and spacer onto row 5 like the real template.
        ws.delete_rows(2, ws.max_row)
        ws["B2"] = "Savings % of Gross"
        ws["B3"] = "Savings % of Net"
        header_row = 4
        for col, value in enumerate(
            ["Mo", "Date", "Account", "Description", "Amount ($)", "Label", "Notes"], start=1
        ):
            ws.cell(row=header_row, column=col, value=value)

        first_data_row = 6
        last_data_row = first_data_row + capacity - 1
        for r in range(first_data_row, last_data_row + 1):
            ws.cell(row=r, column=1, value=f"=MONTH(B{r})")

    wb.save(path)


@pytest.fixture
def tiny_workbook_path(tmp_path) -> Path:
    path = tmp_path / "tiny_workbook.xlsx"
    build_tiny_workbook(path, capacity=5)
    return path


@pytest.fixture
def fixtures_dir() -> Path:
    return Path(__file__).parent / "fixtures"
